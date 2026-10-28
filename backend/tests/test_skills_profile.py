"""Prompt A — resume-to-skills profile: parser fixtures, taxonomy, confirm endpoint, match scores."""
import base64
import io
import os

os.environ.setdefault("DB_NAME", "skills_test_db")

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
import jwt

from bson import ObjectId

from core.db import db
from core.security import JWT_SECRET, JWT_ALGORITHM
from routers import users as users_router
from services.recommendation_engine import calculate_match_score
from services.skills import taxonomy, resume_parser

def make_token(sub: str, role: str = "student") -> str:
    return jwt.encode({"sub": sub, "role": role, "type": "access"}, JWT_SECRET, JWT_ALGORITHM)


def _to_data_url(content: bytes, mime: str) -> str:
    return f"data:{mime};base64,{base64.b64encode(content).decode('ascii')}"


# ─── fixture builders (in-test, no binary assets committed) ─────────────────

def build_text_pdf() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    lines = [
        "ALEX MORA",
        "Software Engineer / Python, React, Docker, Kubernetes, MongoDB",
        "",
        "EXPERIENCE",
        "Software Engineer at Zeta Labs, 2024 - 2026",
        "Built Python APIs with FastAPI; shipped a React dashboard; ran Kubernetes clusters in Docker.",
        "",
        "EDUCATION",
        "B.Tech CSE at ABC University, 2024",
        "",
        "SKILLS",
        "Python, React, FastAPI, Docker, Kubernetes, MongoDB",
    ]
    for line in lines:
        pdf.cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


def build_text_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Intern at Acme — built REST APIs with Python and deployed Docker containers.")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, Django, Docker, Git")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_image_only_pdf() -> bytes:
    from PIL import Image

    img = Image.new("RGB", (400, 300), color=(250, 250, 250))
    buf = io.BytesIO()
    img.save(buf, format="PDF")
    return buf.getvalue()


def sample_user(email: str = "alex@example.com") -> dict:
    return {
        "email": email,
        "first_name": "Alex",
        "last_name": "Mora",
        "role": "student",
        "student_info": {"career_goal": "Backend Developer"},
        "skills": ["Python", "React"],
        "skill_review": {
            "status": "pending",
            "resume_id": "res-1",
            "proposed_skills": [
                {"id": "python", "name": "Python", "category": "general", "confidence": 0.99},
                {"id": "react", "name": "React", "category": "frontend", "confidence": 0.99},
            ],
        },
        "skill_profile": [
            {"id": "python", "name": "Python", "category": "general",
             "source": "resume", "confidence": 0.99, "confirmed_at": "2026-10-27T00:00:00Z"},
            {"id": "react", "name": "React", "category": "frontend",
             "source": "resume", "confidence": 0.99, "confirmed_at": "2026-10-27T00:00:00Z"},
        ],
    }


@pytest.fixture
async def app():
    application = FastAPI()
    application.include_router(users_router.router, prefix="/api")
    yield application


@pytest.fixture
async def client(app):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


# ─── 1. resume parser on 3 fixtures ─────────────────────────────────────────

@pytest.mark.anyio
async def test_parse_text_pdf():
    result = await resume_parser.parse_resume(_to_data_url(build_text_pdf(), "application/pdf"), use_llm=False)
    assert result["text_found"] is True
    names = {s["name"] for s in result["skills"]}
    assert "Python" in names and "Docker" in names and "Kubernetes" in names
    assert result["sha256"]
    assert len(result["text_preview"]) <= resume_parser.MAX_TEXT_BYTES


@pytest.mark.anyio
async def test_parse_docx():
    result = await resume_parser.parse_resume(_to_data_url(build_text_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), use_llm=False)
    assert result["text_found"] is True
    names = {s["name"] for s in result["skills"]}
    assert "Python" in names and "Docker" in names


@pytest.mark.anyio
async def test_parse_image_only_pdf_graceful():
    result = await resume_parser.parse_resume(_to_data_url(build_image_only_pdf(), "application/pdf"), use_llm=False)
    assert result["text_found"] is False
    assert "no text" in result["no_text_reason"].lower()
    assert result["skills"] == []
    # sha256 of the bytes is still computed even when no text layer exists,
    # so unchanged re-uploads are not re-parsed.
    assert len(result["sha256"]) == 64


@pytest.mark.anyio
async def test_sha256_skip_signal():
    pdf_bytes = build_text_pdf()
    h1 = resume_parser.document_sha256(pdf_bytes)
    h2 = resume_parser.document_sha256(pdf_bytes)
    assert h1 == h2 and len(h1) == 64


# ─── 2. taxonomy normalization ───────────────────────────────────────────────

@pytest.mark.anyio
async def test_taxonomy_normalization():
    await taxonomy.load_taxonomy()
    assert taxonomy.normalize("ReactJS")["name"] == "React"
    assert taxonomy.normalize("k8s")["name"] == "Kubernetes"
    assert taxonomy.normalize("docker")["name"] == "Docker"
    assert taxonomy.normalize("zzzqqqqxxx") is None

    hits = taxonomy.extract("I know Python and Kubernetes")
    names = {s["name"] for s in hits}
    assert {"Python", "Kubernetes"} <= names

    sugg = taxonomy.suggest("py")
    assert sugg and sugg[0]["name"] == "Python"


# ─── 3. confirm endpoint ─────────────────────────────────────────────────────

@pytest.mark.anyio
async def test_confirm_endpoint_writes_profile_and_skills(client):
    user_oid = ObjectId()
    token = make_token(str(user_oid))
    await db.users.insert_one({**sample_user(), "_id": user_oid})

    r = await client.post(
        "/api/users/me/skills/confirm",
        json={"add": ["python", "react"], "remove": []},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["confirmed"] == ["python", "react"]
    assert "Python" in body["skills"] and "React" in body["skills"]

    doc = await db.users.find_one({"_id": user_oid})
    assert "Python" in doc["skills"] and "React" in doc["skills"]
    assert all(s.get("source") == "resume" for s in doc["skill_profile"])
    assert doc["skill_review"]["status"] == "confirmed"

    # ids outside the current proposal are rejected
    r2 = await client.post(
        "/api/users/me/skills/confirm",
        json={"add": ["nonexistent-skill-xyz"], "remove": []},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r2.status_code == 400
    assert "not in current resume proposal" in r2.json()["detail"]


@pytest.mark.anyio
async def test_skills_summary_endpoint(client):
    user_oid = ObjectId()
    token = make_token(str(user_oid))
    await db.users.insert_one({**sample_user(), "_id": user_oid})

    r = await client.get("/api/users/me/skills/summary",
                        headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert "Python" in body["skills"]
    assert {p["id"] for p in body["profile"]} == {"python", "react"}
    assert all(p.get("source") == "resume" for p in body["profile"])
    assert body["review"]["status"] == "pending"
    assert {s["id"] for s in body["review"]["proposed_skills"]} == {"python", "react"}

    # blank user — no review, no profile
    blank_oid = ObjectId()
    blank_token = make_token(str(blank_oid))
    await db.users.insert_one({"_id": blank_oid, "email": f"b_{blank_oid}@test.com", "role": "student"})
    r2 = await client.get("/api/users/me/skills/summary",
                         headers={"Authorization": f"Bearer {blank_token}"})
    assert r2.status_code == 200
    b2 = r2.json()
    assert b2["skills"] == [] and b2["profile"] == [] and b2["review"] is None


@pytest.mark.anyio
async def test_suggest_endpoint(client):
    user_oid = ObjectId()
    token = make_token(str(user_oid))
    await db.users.insert_one({"_id": user_oid, "email": f"sugg_{user_oid}@test.com", "role": "student"})

    r = await client.get("/api/skills/suggest", params={"q": "py"},
                         headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200
    assert r.json()["suggestions"][0]["name"] == "Python"

    r2 = await client.get("/api/skills/suggest", params={"q": "k"},
                         headers={"Authorization": f"Bearer {token}"})
    assert r2.status_code == 400


# ─── 4. differentiated match scores ──────────────────────────────────────────

@pytest.mark.anyio
async def test_match_scores_differ_and_none_without_skills(client):
    student_oid = ObjectId()
    blank_oid = ObjectId()
    student_token = make_token(str(student_oid))
    blank_token = make_token(str(blank_oid))
    opp_id = f"opp-{student_oid}"

    user_doc = sample_user()
    await db.users.insert_one({**user_doc, "_id": student_oid})
    await db.users.insert_one({
        "_id": blank_oid, "email": f"blank_{blank_oid}@test.com",
        "first_name": "No", "last_name": "Skills", "role": "student",
        "skills": [], "skill_profile": [],
    })
    await db.internships.insert_one({
        "id": opp_id,
        "type": "internship",
        "title": "Backend Developer",
        "description": "Python APIs, Docker, Kubernetes",
        "tags": ["python", "docker"],
        "skill_ids": ["python", "react", "docker", "kubernetes"],
        "domain": "software",
    })

    opportunity = await db.internships.find_one({"id": opp_id}, {"_id": 0})

    r = await client.get("/api/opportunities",
                        headers={"Authorization": f"Bearer {student_token}"})
    assert r.status_code == 200, r.text
    mine = next(i for i in r.json()["items"] if i.get("id") == opp_id)
    assert mine["match_score"] is not None

    expected = calculate_match_score(opportunity, user_doc)
    blank = calculate_match_score(opportunity, {"email": "blank@example.com", "skill_profile": []})
    assert blank is None
    assert expected != blank
    assert mine["match_score"] == expected

    r2 = await client.get("/api/opportunities",
                         headers={"Authorization": f"Bearer {blank_token}"})
    assert r2.status_code == 200
    blank_item = next(i for i in r2.json()["items"] if i.get("id") == opp_id)
    assert blank_item["match_score"] is None
