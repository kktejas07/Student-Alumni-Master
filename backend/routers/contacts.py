"""
Contacts — FastAPI router for managing contact lists and importing contacts.

Endpoints:
  GET    /api/admin/contacts              list all contacts (paginated, searchable)
  POST   /api/admin/contacts              add a single contact
  POST   /api/admin/contacts/import       import contacts from CSV
  PUT    /api/admin/contacts/{id}         update a contact
  DELETE /api/admin/contacts/{id}         delete a contact
  DELETE /api/admin/contacts/bulk         bulk delete contacts
  
  GET    /api/admin/contact-lists         list all contact lists
  POST   /api/admin/contact-lists         create a contact list
  PUT    /api/admin/contact-lists/{id}    update a contact list
  DELETE /api/admin/contact-lists/{id}    delete a contact list
  POST   /api/admin/contact-lists/{id}/add    add contacts to a list
  POST   /api/admin/contact-lists/{id}/remove remove contacts from a list
"""
from __future__ import annotations
import csv
import io
import secrets
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Query
from pydantic import BaseModel, Field, field_validator
from bson import ObjectId
from core.deps import require_role

router = APIRouter(prefix="/api/admin", tags=["contacts"])

# ── Pydantic schemas ──────────────────────────────────────────────────────────

class ContactCreate(BaseModel):
    name: str = ""
    email: str = ""
    phone: str = ""
    source: str = "manual"  # manual, csv, import, echo, newsletter
    tags: List[str] = []
    lists: List[str] = []
    notes: str = ""
    # Extended fields from Excel
    first_name: str = ""
    last_name: str = ""
    location: str = ""
    zip_code: str = ""
    headline: str = ""
    current_title: str = ""
    current_company: str = ""
    position_start_date: str = ""
    education_degree: str = ""
    education_institution: str = ""
    profile_url: str = ""

    @field_validator("email")
    @classmethod
    def _validate_email(cls, v):
        import re
        if not v or not v.strip():
            return ""
        email_regex = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
        if not re.match(email_regex, v.strip().lower()):
            raise ValueError("Invalid email address")
        return v.strip().lower()


class ContactUpdate(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    tags: Optional[List[str]] = None
    lists: Optional[List[str]] = None
    notes: Optional[str] = None
    subscribed: Optional[bool] = None
    # Extended fields
    first_name: Optional[str] = None
    last_name: Optional[str] = None
    location: Optional[str] = None
    zip_code: Optional[str] = None
    headline: Optional[str] = None
    current_title: Optional[str] = None
    current_company: Optional[str] = None
    position_start_date: Optional[str] = None
    education_degree: Optional[str] = None
    education_institution: Optional[str] = None
    profile_url: Optional[str] = None


class ContactListCreate(BaseModel):
    name: str
    description: str = ""
    color: str = "#7C3AED"


# ── Helpers ───────────────────────────────────────────────────────────────────

def _ser(doc: Optional[dict]) -> dict:
    """Serialize MongoDB document for JSON response."""
    if doc is None:
        return {}
    out = {}
    for k, v in doc.items():
        if isinstance(v, ObjectId):
            out[k] = str(v)
        elif isinstance(v, list):
            out[k] = [_ser(i) if isinstance(i, dict) else (str(i) if isinstance(i, ObjectId) else i) for i in v]
        elif isinstance(v, dict):
            out[k] = _ser(v)
        elif isinstance(v, datetime):
            out[k] = v.isoformat()
        else:
            out[k] = v
    return out


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _get_db(request: Request):
    from db import db
    return db


async def _require_admin(user: dict = Depends(require_role("admin", "super_admin"))):
    """Require an authenticated administrator for contact management."""
    return user


# ── Contact CRUD ──────────────────────────────────────────────────────────────

@router.get("/contacts")
async def list_contacts(
    request: Request,
    q: str = "",
    source: str = "",
    role: str = "",
    list_id: str = "",
    subscribed: str = "",
    include_registered: str = "true",
    page: int = Query(1, ge=1),
    per_page: int = Query(50, ge=1, le=200),
    user=Depends(_require_admin),
):
    """List contacts with search, role filters, source filters, and pagination.
    
    By default includes both imported contacts AND registered users.
    Set include_registered=false to only show imported contacts.
    """
    db = _get_db(request)
    
    items = []
    
    # Fetch imported contacts
    contact_query: Dict[str, Any] = {}
    if q:
        contact_query["$or"] = [
            {"name": {"$regex": q, "$options": "i"}},
            {"email": {"$regex": q, "$options": "i"}},
            {"phone": {"$regex": q, "$options": "i"}},
            {"current_company": {"$regex": q, "$options": "i"}},
            {"location": {"$regex": q, "$options": "i"}},
        ]
    if source:
        contact_query["source"] = source
    if list_id:
        contact_query["lists"] = list_id
    if subscribed:
        contact_query["subscribed"] = subscribed.lower() == "true"
    if role and role.lower() not in ("all", "registered"):
        r_lower = role.lower()
        if r_lower in ("prospect", "prospects"):
            contact_query["is_registered_user"] = {"$ne": True}
        else:
            contact_query["$or"] = [
                {"tags": {"$regex": r_lower, "$options": "i"}},
                {"role": {"$regex": r_lower, "$options": "i"}},
            ]
    
    imported_cursor = db.contacts.find(contact_query).sort("created_at", -1)
    imported_items = []
    async for doc in imported_cursor:
        item = _ser(doc)
        if item.get("phone"):
            from routers.admin.imports import format_e164_phone
            item["phone"] = format_e164_phone(item["phone"], item.get("location", ""))
        imported_items.append(item)
        
    # Also fetch registered users if requested
    registered_users = []
    if include_registered.lower() == "true" and (not source or source in ("registered", "all")):
        user_query: Dict[str, Any] = {"email": {"$exists": True, "$ne": ""}}
        if q:
            user_query["$or"] = [
                {"full_name": {"$regex": q, "$options": "i"}},
                {"email": {"$regex": q, "$options": "i"}},
                {"company": {"$regex": q, "$options": "i"}},
            ]
        if role and role.lower() not in ("all", "registered", "prospect", "prospects"):
            user_query["role"] = {"$regex": role.lower(), "$options": "i"}
        elif role and role.lower() in ("prospect", "prospects"):
            user_query["email"] = "IMPOSSIBLE_USER_MATCH"
            
        cursor_users = db.users.find(user_query).sort("created_at", -1).limit(500)
        async for doc in cursor_users:
            email = (doc.get("email") or "").strip().lower()
            if not email:
                continue
            loc = doc.get("location") or doc.get("city") or ""
            raw_p = doc.get("phone") or ""
            from routers.admin.imports import format_e164_phone
            formatted_p = format_e164_phone(raw_p, loc)
            registered_users.append({
                "_id": str(doc.get("_id", "")),
                "name": doc.get("full_name") or email.split("@")[0].title(),
                "email": email,
                "phone": formatted_p,
                "source": "registered",
                "tags": [doc.get("role", "student")],
                "role": doc.get("role", "student"),
                "lists": [],
                "notes": "",
                "subscribed": True,
                "created_at": doc.get("created_at", ""),
                "updated_at": doc.get("last_login") or doc.get("created_at", ""),
                "current_company": doc.get("company") or doc.get("current_company") or "",
                "current_title": doc.get("designation") or doc.get("title") or "",
                "location": loc,
                "education_institution": doc.get("college_name") or doc.get("college") or "",
                "is_registered_user": True,
            })
    
    # Merge and deduplicate by email
    merged_map = {}
    for c in imported_items:
        em = (c.get("email") or "").lower()
        if em:
            merged_map[em] = c
            
    for u in registered_users:
        em = u["email"].lower()
        if em in merged_map:
            merged_map[em]["is_registered_user"] = True
            if not merged_map[em].get("current_company"):
                merged_map[em]["current_company"] = u.get("current_company", "")
        else:
            merged_map[em] = u

    all_contacts = list(merged_map.values())
    total = len(all_contacts)
    
    skip = (page - 1) * per_page
    items = all_contacts[skip:skip + per_page]
    
    return {
        "items": items,
        "total": total,
        "imported_count": len(imported_items),
        "registered_count": len(registered_users),
        "page": page,
        "per_page": per_page,
        "total_pages": (total + per_page - 1) // per_page if total > 0 else 1,
    }


@router.post("/contacts")
async def create_contact(payload: ContactCreate, request: Request, user=Depends(_require_admin)):
    """Add a single contact."""
    db = _get_db(request)
    
    # Check for duplicate email
    if payload.email:
        existing = await db.contacts.find_one({"email": payload.email.lower()})
        if existing:
            raise HTTPException(status_code=400, detail="Contact with this email already exists")
    
    contact = {
        "name": payload.name.strip(),
        "email": payload.email.lower().strip(),
        "phone": payload.phone.strip(),
        "source": payload.source,
        "tags": payload.tags,
        "lists": payload.lists,
        "notes": payload.notes,
        "subscribed": True,
        "created_at": _now(),
        "updated_at": _now(),
    }
    
    result = await db.contacts.insert_one(dict(contact))
    contact["_id"] = str(result.inserted_id)
    
    return {"ok": True, "contact": _ser(contact)}


@router.put("/contacts/{contact_id}")
async def update_contact(contact_id: str, payload: ContactUpdate, request: Request, user=Depends(_require_admin)):
    """Update a contact."""
    db = _get_db(request)
    
    update_data = {k: v for k, v in payload.model_dump().items() if v is not None}
    update_data["updated_at"] = _now()
    
    result = await db.contacts.update_one(
        {"_id": ObjectId(contact_id)},
        {"$set": update_data}
    )
    
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Contact not found")
    
    return {"ok": True}


@router.delete("/contacts/bulk")
async def bulk_delete_contacts(payload: Dict[str, List[str]], request: Request, user=Depends(_require_admin)):
    """Bulk delete contacts by IDs."""
    db = _get_db(request)
    ids = payload.get("ids", [])
    
    if not ids:
        raise HTTPException(status_code=400, detail="No contact IDs provided")
    
    result = await db.contacts.delete_many({"_id": {"$in": [ObjectId(id) for id in ids]}})
    
    return {"ok": True, "deleted": result.deleted_count}


# ── Cleanup: Delete Contacts Collection ────────────────────────────────────────

@router.delete("/contacts/cleanup")
async def cleanup_contacts(request: Request, user=Depends(_require_admin)):
    """Delete the contacts collection after migration.
    
    WARNING: This is irreversible. Only use after confirming migration is complete.
    """
    db = _get_db(request)
    
    result = await db.contacts.delete_many({})
    
    return {
        "ok": True,
        "deleted": result.deleted_count,
        "message": f"Deleted {result.deleted_count} contacts. Newsletter will now only use users collection."
    }


@router.delete("/contacts/{contact_id}")
async def delete_contact(contact_id: str, request: Request, user=Depends(_require_admin)):
    """Delete a contact."""
    db = _get_db(request)
    
    result = await db.contacts.delete_one({"_id": ObjectId(contact_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Contact not found")
    
    return {"ok": True}


# ── CSV Import ────────────────────────────────────────────────────────────────

@router.post("/contacts/import")
async def import_contacts(
    request: Request,
    file: UploadFile = File(...),
    source: str = "csv",
    tags: str = "",
    user=Depends(_require_admin),
):
    """Import contacts from a CSV or XLSX file.
    
    Expected columns: name, email, phone (email is required)
    Supports both .csv and .xlsx formats.
    """
    db = _get_db(request)
    
    filename = file.filename or ""
    is_csv = filename.lower().endswith('.csv')
    is_xlsx = filename.lower().endswith('.xlsx') or filename.lower().endswith('.xls')
    
    if not is_csv and not is_xlsx:
        raise HTTPException(status_code=400, detail="File must be a CSV or XLSX file")
    
    content = await file.read()
    rows = []
    
    if is_csv:
        # Parse CSV
        try:
            text = content.decode('utf-8')
        except UnicodeDecodeError:
            try:
                text = content.decode('latin-1')
            except:
                raise HTTPException(status_code=400, detail="Could not decode file. Please use UTF-8 encoding.")
        
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append(row)
    else:
        # Parse XLSX
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        
        if ws is None:
            raise HTTPException(status_code=400, detail="Excel file is empty")
        
        # Get headers from first row
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            headers.append(str(cell.value or "").strip().lower())
        
        # Parse data rows
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for i, value in enumerate(row):
                if i < len(headers) and headers[i]:
                    row_dict[headers[i]] = str(value or "").strip()
            rows.append(row_dict)
        
        wb.close()
    
    # Process rows
    imported = 0
    updated = 0
    skipped = 0
    errors = []
    tag_list = [t.strip() for t in tags.split(",") if t.strip()]
    
    import re
    email_regex = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    
    def get_field(row, *keys, default=""):
        """Get value from row using multiple possible column names."""
        for key in keys:
            val = row.get(key, "").strip()
            if val:
                return val
        return default
    
    for i, row in enumerate(rows, start=2):
        # Get email (required)
        email = get_field(row, 
            "email", "Email", "EMAIL", "email_address", "Email Address",
            "e-mail", "E-Mail", "email address"
        ).lower()
        
        if not email:
            skipped += 1
            errors.append(f"Row {i}: No email found")
            continue
        
        # Validate email
        if not re.match(email_regex, email):
            skipped += 1
            errors.append(f"Row {i}: Invalid email '{email}'")
            continue
        
        # Get name fields
        first_name = get_field(row, "first_name", "First Name", "first name", "firstname")
        last_name = get_field(row, "last_name", "Last Name", "last name", "lastname")
        
        # Combine first + last name, or use name/full_name
        name = (
            f"{first_name} {last_name}".strip() if first_name or last_name else
            get_field(row, "name", "Name", "full_name", "Full Name", "full name") or
            email.split("@")[0].title()
        )
        
        # Get other fields
        phone = get_field(row, 
            "phone", "Phone", "phone_number", "Phone Number", 
            "phone number", "mobile", "Mobile", "mobile_number"
        )
        location = get_field(row, "location", "Location", "general_location", "General Location", "city", "City")
        zip_code = get_field(row, "zip_code", "ZIP Code", "zip", "ZIP", "zipcode", "postal_code")
        headline = get_field(row, "headline", "Headline", "title", "Title")
        current_title = get_field(row, "current_title", "Current Title", "job_title", "Job Title", "position")
        current_company = get_field(row, "current_company", "Current Company", "company", "Company", "organization")
        position_start_date = get_field(row, "position_start_date", "Position Start Date", "start_date", "Start Date")
        education_degree = get_field(row, "education_degree", "Education Degree", "degree", "Degree")
        education_institution = get_field(row, "education_institution", "Education Institution", "school", "School", "university")
        profile_url = get_field(row, "profile_url", "Profile URL", "linkedin", "LinkedIn", "profile", "url")
        
        # Check for existing contact OR user
        existing_contact = await db.contacts.find_one({"email": email})
        existing_user = await db.users.find_one({"email": email})
        
        # Extract role / category column
        role_col = get_field(row, "role", "Role", "category", "Category", "type", "Type", "user_type", "User Type").lower()
        row_tags = list(tag_list)
        if role_col and role_col not in row_tags:
            row_tags.append(role_col)
        
        contact_data = {
            "name": name,
            "email": email,
            "phone": phone,
            "first_name": first_name,
            "last_name": last_name,
            "location": location,
            "zip_code": zip_code,
            "headline": headline,
            "current_title": current_title,
            "current_company": current_company,
            "position_start_date": position_start_date,
            "education_degree": education_degree,
            "education_institution": education_institution,
            "profile_url": profile_url,
            "source": source,
            "tags": row_tags,
            "is_registered_user": bool(existing_user),
            "updated_at": _now(),
        }
        
        if existing_contact:
            # Update existing contact with new non-empty data (merge)
            update_data = {k: v for k, v in contact_data.items() if v}
            update_data["updated_at"] = _now()
            existing_tags = set(existing_contact.get("tags") or [])
            existing_tags.update(row_tags)
            update_data["tags"] = list(existing_tags)
            
            await db.contacts.update_one(
                {"_id": existing_contact["_id"]},
                {"$set": update_data}
            )
            updated += 1
        else:
            # Create new contact
            contact_data["lists"] = []
            contact_data["notes"] = ""
            contact_data["subscribed"] = True
            contact_data["created_at"] = _now()
            await db.contacts.insert_one(dict(contact_data))
            imported += 1
    
    return {
        "ok": True,
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:10],
        "total_rows": imported + updated + skipped,
        "file_type": "xlsx" if is_xlsx else "csv",
    }


# ── Contact Lists ─────────────────────────────────────────────────────────────

@router.get("/contact-lists")
async def list_contact_lists(request: Request, user=Depends(_require_admin)):
    """List all contact lists with contact counts."""
    db = _get_db(request)
    
    cursor = db.contact_lists.find().sort("name", 1)
    lists = []
    async for doc in cursor:
        list_doc = _ser(doc)
        # Get contact count for this list
        count = await db.contacts.count_documents({"lists": doc.get("_id") or doc.get("name")})
        list_doc["contact_count"] = count
        lists.append(list_doc)
    
    return {"items": lists}


@router.post("/contact-lists")
async def create_contact_list(payload: ContactListCreate, request: Request, user=Depends(_require_admin)):
    """Create a new contact list."""
    db = _get_db(request)
    
    # Check for duplicate name
    existing = await db.contact_lists.find_one({"name": payload.name})
    if existing:
        raise HTTPException(status_code=400, detail="A list with this name already exists")
    
    contact_list = {
        "name": payload.name.strip(),
        "description": payload.description,
        "color": payload.color,
        "created_at": _now(),
        "updated_at": _now(),
    }
    
    result = await db.contact_lists.insert_one(dict(contact_list))
    contact_list["_id"] = str(result.inserted_id)
    
    return {"ok": True, "list": _ser(contact_list)}


@router.put("/contact-lists/{list_id}")
async def update_contact_list(list_id: str, payload: ContactListCreate, request: Request, user=Depends(_require_admin)):
    """Update a contact list."""
    db = _get_db(request)
    
    result = await db.contact_lists.update_one(
        {"_id": ObjectId(list_id)},
        {"$set": {
            "name": payload.name,
            "description": payload.description,
            "color": payload.color,
            "updated_at": _now(),
        }}
    )
    
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="List not found")
    
    return {"ok": True}


@router.delete("/contact-lists/{list_id}")
async def delete_contact_list(list_id: str, request: Request, user=Depends(_require_admin)):
    """Delete a contact list."""
    db = _get_db(request)
    
    result = await db.contact_lists.delete_one({"_id": ObjectId(list_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="List not found")
    
    return {"ok": True}


@router.post("/contact-lists/{list_id}/add")
async def add_contacts_to_list(list_id: str, payload: Dict[str, List[str]], request: Request, user=Depends(_require_admin)):
    """Add contacts to a list."""
    db = _get_db(request)
    
    contact_ids = payload.get("contact_ids", [])
    if not contact_ids:
        raise HTTPException(status_code=400, detail="No contact IDs provided")
    
    # Get list name
    contact_list = await db.contact_lists.find_one({"_id": ObjectId(list_id)})
    if not contact_list:
        raise HTTPException(status_code=404, detail="List not found")
    
    list_name = contact_list.get("name", "")
    
    # Add list to contacts
    result = await db.contacts.update_many(
        {"_id": {"$in": [ObjectId(id) for id in contact_ids]}},
        {"$addToSet": {"lists": list_name}}
    )
    
    return {"ok": True, "updated": result.modified_count}


@router.post("/contact-lists/{list_id}/remove")
async def remove_contacts_from_list(list_id: str, payload: Dict[str, List[str]], request: Request, user=Depends(_require_admin)):
    """Remove contacts from a list."""
    db = _get_db(request)
    
    contact_ids = payload.get("contact_ids", [])
    if not contact_ids:
        raise HTTPException(status_code=400, detail="No contact IDs provided")
    
    # Get list name
    contact_list = await db.contact_lists.find_one({"_id": ObjectId(list_id)})
    if not contact_list:
        raise HTTPException(status_code=404, detail="List not found")
    
    list_name = contact_list.get("name", "")
    
    # Remove list from contacts
    result = await db.contacts.update_many(
        {"_id": {"$in": [ObjectId(id) for id in contact_ids]}},
        {"$pull": {"lists": list_name}}
    )
    
    return {"ok": True, "updated": result.modified_count}


# ── Stats ─────────────────────────────────────────────────────────────────────

# ── Stats ─────────────────────────────────────────────────────────────────────

@router.get("/contacts/stats")
async def contact_stats(request: Request, user=Depends(_require_admin)):
    """Get comprehensive contact and registered user statistics."""
    db = _get_db(request)
    
    total_contacts = await db.contacts.count_documents({})
    total_users = await db.users.count_documents({})
    
    subscribed = await db.contacts.count_documents({"subscribed": True})
    unsubscribed = await db.contacts.count_documents({"subscribed": False})
    
    # Role breakdown across users & contacts
    student_users = await db.users.count_documents({"role": "student"})
    mentor_users = await db.users.count_documents({"role": "mentor"})
    alumni_users = await db.users.count_documents({"role": "alumni"})
    college_users = await db.users.count_documents({"role": {"$in": ["college", "admin", "superadmin"]}})
    
    student_contacts = await db.contacts.count_documents({"tags": {"$regex": "student", "$options": "i"}})
    mentor_contacts = await db.contacts.count_documents({"tags": {"$regex": "mentor", "$options": "i"}})
    alumni_contacts = await db.contacts.count_documents({"tags": {"$regex": "alumni", "$options": "i"}})
    college_contacts = await db.contacts.count_documents({"tags": {"$regex": "college|admin", "$options": "i"}})

    prospects_count = await db.contacts.count_documents({"is_registered_user": {"$ne": True}})

    # Source breakdown
    pipeline = [
        {"$group": {"_id": "$source", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
    ]
    sources = {"registered": total_users}
    async for doc in db.contacts.aggregate(pipeline):
        if doc["_id"]:
            sources[doc["_id"]] = doc["count"]
    
    lists = await db.contact_lists.count_documents({})
    
    return {
        "total": total_contacts + total_users,
        "registered_count": total_users,
        "imported_count": total_contacts,
        "prospects_count": prospects_count,
        "students_count": student_users + student_contacts,
        "mentors_count": mentor_users + mentor_contacts,
        "alumni_count": alumni_users + alumni_contacts,
        "college_count": college_users + college_contacts,
        "subscribed": subscribed,
        "unsubscribed": unsubscribed,
        "sources": sources,
        "lists": lists,
    }


# ── Export & Share Emails ─────────────────────────────────────────────────────

@router.get("/contacts/export")
async def export_contacts(
    request: Request,
    q: str = "",
    source: str = "",
    role: str = "",
    format: str = "json",
    user=Depends(_require_admin),
):
    """Export contacts or extract email address list for sending broadcasts/sharing."""
    res = await list_contacts(request, q=q, source=source, role=role, page=1, per_page=10000, user=user)
    items = res.get("items", [])
    
    emails = [item["email"] for item in items if item.get("email")]
    email_string = ", ".join(emails)
    
    if format == "emails":
        return {"count": len(emails), "email_string": email_string, "emails": emails}
    elif format == "csv":
        from fastapi.responses import Response
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=["name", "email", "phone", "role", "source", "current_company", "current_title", "location"])
        writer.writeheader()
        for item in items:
            writer.writerow({
                "name": item.get("name", ""),
                "email": item.get("email", ""),
                "phone": item.get("phone", ""),
                "role": ", ".join(item.get("tags") or []) if item.get("tags") else item.get("role", ""),
                "source": item.get("source", ""),
                "current_company": item.get("current_company", ""),
                "current_title": item.get("current_title", ""),
                "location": item.get("location", ""),
            })
        return Response(content=output.getvalue(), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=contacts_export.csv"})
    
    return {"items": items, "count": len(items), "emails": emails, "email_string": email_string}


# ── Migration: Contacts → Users ───────────────────────────────────────────────

@router.post("/contacts/migrate-to-users")
async def migrate_contacts_to_users(request: Request, user=Depends(_require_admin)):
    """Migrate all contacts to the users collection.
    
    - Creates user accounts for contacts that don't have one
    - Skips duplicates (email already exists in users)
    - Makes email mandatory
    - Returns migration stats
    """
    db = _get_db(request)
    
    migrated = 0
    skipped = 0
    errors = []
    
    # Get all contacts
    cursor = db.contacts.find({})
    async for contact in cursor:
        email = (contact.get("email") or "").strip().lower()
        
        # Skip if no email
        if not email:
            skipped += 1
            continue
        
        # Check if user already exists
        existing_user = await db.users.find_one({"email": email})
        if existing_user:
            skipped += 1
            continue
        
        # Create user from contact
        try:
            from core.security import hash_password
            
            # Generate temp password
            import secrets as _secrets
            temp_password = _secrets.token_urlsafe(12)
            
            # Determine role from tags or default to student
            tags = contact.get("tags", [])
            role = "student"
            if "mentor" in tags:
                role = "mentor"
            elif "alumni" in tags:
                role = "alumni"
            
            user_doc = {
                "email": email,
                "password_hash": hash_password(temp_password),
                "full_name": contact.get("name") or email.split("@")[0].title(),
                "first_name": contact.get("first_name") or "",
                "last_name": contact.get("last_name") or "",
                "phone": contact.get("phone") or "",
                "role": role,
                "category": role,
                "company": contact.get("current_company") or "",
                "designation": contact.get("current_title") or "",
                "location": contact.get("location") or "",
                "college": contact.get("education_institution") or "",
                "degree": contact.get("education_degree") or "",
                "linkedin_url": contact.get("profile_url") or "",
                "created_at": datetime.now(timezone.utc),
                "created_by": "migration",
                "approved": True,
            }
            
            await db.users.insert_one(user_doc)
            migrated += 1
            
        except Exception as e:
            errors.append(f"Failed to migrate {email}: {str(e)}")
    
    return {
        "ok": True,
        "migrated": migrated,
        "skipped": skipped,
        "errors": errors[:10],
        "message": f"Migrated {migrated} contacts to users. {skipped} skipped (duplicate or no email)."
    }
