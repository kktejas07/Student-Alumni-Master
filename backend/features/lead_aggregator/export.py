"""Enhancement #3 — Lead/Brand/Claim Export as CSV and Excel.

Provides streaming file downloads for the Lead Aggregator dashboard.
"""
from __future__ import annotations

import csv
import io
import json
import logging
from datetime import datetime, timezone
from typing import Any

from fastapi.responses import StreamingResponse

from .database import leads_col, brands_col, claims_col

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

LEAD_COLUMNS = [
    "lead_id", "full_name", "email", "phone", "age", "city", "domain",
    "score", "temperature", "status", "phone_verified", "profile_completeness",
    "matched_brands", "created_at",
]

BRAND_COLUMNS = [
    "brand_id", "name", "domain", "tier", "status", "contact_email",
    "contact_phone", "webhook_url", "lead_quota", "target_cities",
    "target_skills", "target_education_levels", "created_at",
]

CLAIM_COLUMNS = [
    "claim_id", "lead_id", "offer_id", "brand_id", "phone", "email",
    "status", "claimed_at",
]


def _flatten_lead(lead: dict) -> dict:
    flat: dict[str, Any] = {}
    for col in LEAD_COLUMNS:
        val = lead.get(col)
        if col == "matched_brands" and isinstance(val, list):
            flat[col] = ", ".join(val)
        elif col == "created_at" and isinstance(val, datetime):
            flat[col] = val.strftime("%Y-%m-%d %H:%M:%S")
        elif isinstance(val, (dict, list)):
            flat[col] = json.dumps(val, default=str)
        else:
            flat[col] = val if val is not None else ""
    return flat


def _flatten_brand(brand: dict) -> dict:
    flat: dict[str, Any] = {}
    for col in BRAND_COLUMNS:
        val = brand.get(col)
        if isinstance(val, list):
            flat[col] = ", ".join(str(v) for v in val)
        elif isinstance(val, datetime):
            flat[col] = val.strftime("%Y-%m-%d %H:%M:%S")
        elif isinstance(val, dict):
            flat[col] = json.dumps(val, default=str)
        else:
            flat[col] = val if val is not None else ""
    return flat


def _flatten_claim(claim: dict) -> dict:
    flat: dict[str, Any] = {}
    for col in CLAIM_COLUMNS:
        val = claim.get(col)
        if isinstance(val, datetime):
            flat[col] = val.strftime("%Y-%m-%d %H:%M:%S")
        else:
            flat[col] = val if val is not None else ""
    return flat


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


async def _collect_docs(col, query: dict, columns: list[str], flatten_fn, limit: int) -> list[dict]:
    cursor = col.find(query, {"_id": 0}).sort("created_at", -1).limit(limit)
    rows = []
    async for doc in cursor:
        rows.append(flatten_fn(doc))
    return rows


def _build_csv(columns: list[str], rows: list[dict]) -> io.StringIO:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    buf.seek(0)
    return buf


async def export_leads_csv(query: dict, limit: int = 10000) -> StreamingResponse:
    rows = await _collect_docs(leads_col, query, LEAD_COLUMNS, _flatten_lead, limit)
    buf = _build_csv(LEAD_COLUMNS, rows)
    filename = f"leads_export_{_timestamp()}.csv"
    return StreamingResponse(
        buf,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def export_leads_excel(query: dict, limit: int = 10000) -> StreamingResponse:
    if not HAS_OPENPYXL:
        rows = await _collect_docs(leads_col, query, LEAD_COLUMNS, _flatten_lead, limit)
        buf = _build_csv(LEAD_COLUMNS, rows)
        return StreamingResponse(
            buf,
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="leads_export_{_timestamp()}.csv"'},
        )

    rows = await _collect_docs(leads_col, query, LEAD_COLUMNS, _flatten_lead, limit)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    header_align = Alignment(horizontal="center")

    for col_idx, col_name in enumerate(LEAD_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row_idx, row in enumerate(rows, 2):
        for col_idx, col_name in enumerate(LEAD_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

    for col_idx in range(1, len(LEAD_COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"leads_export_{_timestamp()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def export_brands_csv(query: dict, limit: int = 5000) -> StreamingResponse:
    rows = await _collect_docs(brands_col, query, BRAND_COLUMNS, _flatten_brand, limit)
    buf = _build_csv(BRAND_COLUMNS, rows)
    filename = f"brands_export_{_timestamp()}.csv"
    return StreamingResponse(
        buf,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def export_claims_csv(query: dict, limit: int = 10000) -> StreamingResponse:
    rows = await _collect_docs(claims_col, query, CLAIM_COLUMNS, _flatten_claim, limit)
    buf = _build_csv(CLAIM_COLUMNS, rows)
    filename = f"claims_export_{_timestamp()}.csv"
    return StreamingResponse(
        buf,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
